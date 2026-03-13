from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from datetime import datetime

def create_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    
    title.text = "PG Allocator"
    title.text_frame.paragraphs[0].font.size = Pt(44)
    
    subtitle.text = "A Complete PG/Hostel Management Solution\n" + datetime.now().strftime("%B %d, %Y")
    subtitle.text_frame.paragraphs[0].font.size = Pt(24)

def add_slide_with_title_and_content(prs, title_text, content_list):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    content = slide.placeholders[1]
    
    title.text = title_text
    
    for item in content_list:
        p = content.text_frame.add_paragraph()
        p.text = item
        p.space_after = Pt(12)
        p.level = 0

def create_presentation():
    prs = Presentation()
    
    # Slide 1: Title
    create_title_slide(prs)
    
    # Slide 2: Project Overview
    overview = [
        "Full-stack web application for managing PG/Hostel accommodations",
        "Features user authentication, room allocation, and payment tracking",
        "Built with modern web technologies"
    ]
    add_slide_with_title_and_content(prs, "Project Overview", overview)
    
    # Slide 3: Technology Stack
    tech_stack = [
        "Frontend: React.js",
        "Backend: Java",
        "Database: MySQL",
        "Web Server: XAMPP",
        "APIs: RESTful services"
    ]
    add_slide_with_title_and_content(prs, "Technology Stack", tech_stack)
    
    # Slide 4: System Architecture
    architecture = [
        "Frontend Layer: React components and pages",
        "Backend Layer: Java-based server",
        "Data Access Layer: DAO pattern implementation",
        "Database Layer: MySQL database"
    ]
    add_slide_with_title_and_content(prs, "System Architecture", architecture)
    
    # Slide 5: Key Features
    features = [
        "User Authentication",
        "Room Management",
        "Booking System",
        "Payment Tracking",
        "Admin Dashboard",
        "Tenant Management"
    ]
    add_slide_with_title_and_content(prs, "Key Features", features)
    
    # Slide 6: Database Schema
    db_schema = [
        "Users: User accounts and authentication",
        "Bookings: Room reservation details",
        "Rooms: Room information and availability",
        "Payments: Transaction records",
        "PG/Hostel: Property details"
    ]
    add_slide_with_title_and_content(prs, "Database Schema", db_schema)
    
    # Slide 7: Frontend Structure
    frontend = [
        "Components: Reusable UI elements",
        "Pages: Main application views",
        "Contexts: State management",
        "Services: API communication"
    ]
    add_slide_with_title_and_content(prs, "Frontend Structure", frontend)
    
    # Slide 8: Backend Structure
    backend = [
        "Models: Data structures",
        "DAO: Database operations",
        "Utils: Helper functions",
        "Server: Main application server"
    ]
    add_slide_with_title_and_content(prs, "Backend Structure", backend)
    
    # Slide 9: Setup & Installation
    setup = [
        "1. Install XAMPP",
        "2. Set up MySQL database",
        "3. Configure backend",
        "4. Install frontend dependencies",
        "5. Start both frontend and backend servers"
    ]
    add_slide_with_title_and_content(prs, "Setup & Installation", setup)
    
    # Slide 10: Future Enhancements
    future = [
        "Mobile application development",
        "Advanced reporting and analytics",
        "Online payment integration",
        "Review and rating system",
        "Automated notifications"
    ]
    add_slide_with_title_and_content(prs, "Future Enhancements", future)
    
    # Slide 11: Demo
    demo = [
        "Screenshots of key features",
        "Live demo of the application",
        "User flow demonstration"
    ]
    add_slide_with_title_and_content(prs, "Demo", demo)
    
    # Slide 12: Q&A
    qa_slide = prs.slides.add_slide(prs.slide_layouts[5])
    title = qa_slide.shapes.title
    title.text = "Q&A"
    
    # Save the presentation
    prs.save('PG_Allocator_Presentation.pptx')

if __name__ == "__main__":
    create_presentation()
    print("Presentation created successfully: PG_Allocator_Presentation.pptx")
